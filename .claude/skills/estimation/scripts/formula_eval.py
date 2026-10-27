#!/usr/bin/env python3
"""formula_eval.py - a small Excel formula evaluator for the estimation selftest.

It evaluates only the function set that build_workbook.py writes, so the
selftest can check what every formula will produce without Excel installed:

    SUM SUMIFS SUMPRODUCT COUNTIF COUNTIFS COUNTBLANK IF IFERROR AND OR NOT
    ABS MAX MIN ROUND ROUNDUP INDEX MATCH WORKDAY ISNUMBER LEFT

Operators: + - * / & = <> < > <= >= and unary minus. Ranges are flattened to
lists and arithmetic broadcasts element-wise, which is how SUMPRODUCT treats
them in Excel. Real Excel (or LibreOffice) recalculation stays the final check.
"""

import datetime as _dt
import math
import re

EXCEL_EPOCH = _dt.datetime(1899, 12, 30)


class XLError(Exception):
    """An Excel error value such as #N/A or #VALUE!."""


TOKEN_RE = re.compile(
    r"""\s*(?:
    (?P<str>"(?:[^"]|"")*")
  | (?P<ref>(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_\.]*)!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?
        |\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
  | (?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)
  | (?P<func>[A-Z][A-Z0-9\.]*)\s*\(
  | (?P<bool>TRUE|FALSE)\b
  | (?P<op><>|<=|>=|[-+*/&=<>(),])
    )""",
    re.X,
)

CELL_RE = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")


def col_to_num(col):
    """Convert a column letter (A, AB) to a 1-based index."""
    n = 0
    for ch in col:
        n = n * 26 + ord(ch) - 64
    return n


def tokenize(text):
    """Split a formula body (without the leading '=') into tokens."""
    pos, out = 0, []
    text = text.strip()
    while pos < len(text):
        m = TOKEN_RE.match(text, pos)
        if not m or m.end() == pos:
            raise ValueError("cannot tokenize formula at: {0!r}".format(text[pos:]))
        kind = m.lastgroup
        out.append((kind, m.group(kind)))
        pos = m.end()
        while pos < len(text) and text[pos] == " ":
            pos += 1
    return out


class Parser:
    """Recursive-descent parser producing a nested-tuple AST."""

    def __init__(self, tokens):
        self.t = tokens
        self.i = 0

    def peek(self):
        """Return the current token or (None, None)."""
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def take(self):
        """Consume and return the current token."""
        tok = self.peek()
        self.i += 1
        return tok

    def expect(self, value):
        """Consume a specific operator token."""
        tok = self.take()
        if tok[1] != value:
            raise ValueError("expected {0!r}, got {1!r}".format(value, tok))

    def parse(self):
        """Parse a full expression."""
        node = self.comparison()
        if self.i != len(self.t):
            raise ValueError("trailing tokens: {0!r}".format(self.t[self.i:]))
        return node

    def comparison(self):
        """Lowest precedence: = <> < > <= >=."""
        node = self.concat()
        while self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.take()[1]
            node = ("bin", op, node, self.concat())
        return node

    def concat(self):
        """String concatenation with &."""
        node = self.additive()
        while self.peek()[1] == "&":
            self.take()
            node = ("bin", "&", node, self.additive())
        return node

    def additive(self):
        """Addition and subtraction."""
        node = self.term()
        while self.peek()[1] in ("+", "-"):
            op = self.take()[1]
            node = ("bin", op, node, self.term())
        return node

    def term(self):
        """Multiplication and division."""
        node = self.unary()
        while self.peek()[1] in ("*", "/"):
            op = self.take()[1]
            node = ("bin", op, node, self.unary())
        return node

    def unary(self):
        """Unary minus and plus."""
        if self.peek()[1] in ("-", "+"):
            op = self.take()[1]
            inner = self.unary()
            return ("neg", inner) if op == "-" else inner
        return self.primary()

    def primary(self):
        """Literals, references, function calls, parentheses."""
        kind, val = self.take()
        if kind == "num":
            return ("lit", float(val))
        if kind == "str":
            return ("lit", val[1:-1].replace('""', '"'))
        if kind == "bool":
            return ("lit", val == "TRUE")
        if kind == "ref":
            return ("ref", val)
        if kind == "func":
            args = []
            if self.peek()[1] != ")":
                args.append(self.comparison())
                while self.peek()[1] == ",":
                    self.take()
                    args.append(self.comparison())
            self.expect(")")
            return ("call", val, args)
        if val == "(":
            node = self.comparison()
            self.expect(")")
            return node
        raise ValueError("unexpected token {0!r}".format((kind, val)))


def is_num(v):
    """True for int/float but not bool."""
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def to_num(v):
    """Coerce a scalar to a number the way Excel arithmetic does."""
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if is_num(v):
        return float(v)
    try:
        return float(v)
    except (TypeError, ValueError):
        raise XLError("#VALUE!")


def to_text(v):
    """Coerce a scalar to text the way Excel & does."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if is_num(v):
        return str(int(v)) if float(v).is_integer() else "{0:.15g}".format(v)
    return str(v)


def compare(a, b, op):
    """Excel comparison: blanks equal 0 and '', text is case-insensitive."""
    if a is None:
        a = "" if isinstance(b, str) else 0.0
    if b is None:
        b = "" if isinstance(a, str) else 0.0
    if isinstance(a, bool):
        a = 1.0 if a else 0.0
    if isinstance(b, bool):
        b = 1.0 if b else 0.0
    if isinstance(a, str) and isinstance(b, str):
        a, b = a.lower(), b.lower()
    elif isinstance(a, str) != isinstance(b, str):
        # Text always sorts above numbers in Excel.
        a, b = (1, 0) if isinstance(a, str) else (0, 1)
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b,
            "<=": a <= b, ">=": a >= b}[op]


def crit_match(value, crit):
    """SUMIFS/COUNTIF criteria matching for the criteria this skill writes."""
    if is_num(crit):
        return is_num(value) and float(value) == float(crit)
    crit = "" if crit is None else str(crit)
    op = "="
    for cand in ("<=", ">=", "<>", "<", ">", "="):
        if crit.startswith(cand):
            op, crit = cand, crit[len(cand):]
            break
    blank = value is None or value == ""
    if crit == "":
        return blank if op == "=" else (not blank if op == "<>" else False)
    try:
        cnum = float(crit)
    except ValueError:
        cnum = None
    if cnum is not None and is_num(value):
        return compare(float(value), cnum, op)
    if cnum is not None and op not in ("=", "<>"):
        return False
    if blank:
        return op == "<>"
    if is_num(value):
        return op == "<>"
    return compare(str(value), crit, op)


def flatten(v):
    """Flatten nested lists to one list."""
    if isinstance(v, list):
        out = []
        for x in v:
            out.extend(flatten(x))
        return out
    return [v]


def broadcast(fn, a, b):
    """Apply a binary scalar function element-wise over lists."""
    if isinstance(a, list) or isinstance(b, list):
        la = a if isinstance(a, list) else None
        lb = b if isinstance(b, list) else None
        n = len(la) if la is not None else len(lb)
        if la is not None and lb is not None and len(la) != len(lb):
            raise XLError("#VALUE!")
        return [fn(la[i] if la is not None else a, lb[i] if lb is not None else b)
                for i in range(n)]
    return fn(a, b)


def workday(start, days):
    """Excel WORKDAY with no holiday list (Saturday and Sunday skipped)."""
    d = int(math.floor(start))
    days = int(days)
    step = 1 if days >= 0 else -1
    remaining = abs(days)
    while remaining:
        d += step
        if (EXCEL_EPOCH + _dt.timedelta(days=d)).weekday() < 5:
            remaining -= 1
    return float(d)


def roundup(x, digits):
    """Excel ROUNDUP: away from zero."""
    f = 10 ** int(digits)
    return math.copysign(math.ceil(abs(x) * f) / f, x)


class Evaluator:
    """Evaluate formulas in an openpyxl workbook loaded with data_only=False."""

    def __init__(self, wb):
        self.wb = wb
        self.cache = {}
        self.ast_cache = {}

    def cell(self, sheet, coord):
        """Return the evaluated value of one cell."""
        key = (sheet, coord)
        if key in self.cache:
            return self.cache[key]
        raw = self.wb[sheet][coord].value
        if isinstance(raw, str) and raw.startswith("="):
            self.cache[key] = XLError("#CYCLE")
            try:
                val = self.formula(sheet, raw)
            except XLError as err:
                val = err
            self.cache[key] = val
            return val
        if isinstance(raw, _dt.datetime):
            raw = float((raw - EXCEL_EPOCH).days)
        elif isinstance(raw, _dt.date):
            raw = float((_dt.datetime(raw.year, raw.month, raw.day) - EXCEL_EPOCH).days)
        elif isinstance(raw, int) and not isinstance(raw, bool):
            raw = float(raw)
        self.cache[key] = raw
        return raw

    def value(self, sheet, coord):
        """Public accessor: raise on an error value."""
        v = self.cell(sheet, coord)
        if isinstance(v, XLError):
            raise v
        return v

    def formula(self, sheet, text):
        """Parse (cached) and evaluate a formula string."""
        if text not in self.ast_cache:
            self.ast_cache[text] = Parser(tokenize(text[1:])).parse()
        return self.ev(self.ast_cache[text], sheet)

    def ref(self, sheet, text):
        """Resolve a reference to a scalar (single cell) or a list (range)."""
        if "!" in text:
            sname, addr = text.rsplit("!", 1)
            if sname.startswith("'"):
                sname = sname[1:-1].replace("''", "'")
        else:
            sname, addr = sheet, text
        parts = addr.split(":")
        m1 = CELL_RE.fullmatch(parts[0])
        if len(parts) == 1:
            return self.scalar(self.cell(sname, m1.group(1) + m1.group(2)))
        m2 = CELL_RE.fullmatch(parts[1])
        c1, c2 = col_to_num(m1.group(1)), col_to_num(m2.group(1))
        r1, r2 = int(m1.group(2)), int(m2.group(2))
        from openpyxl.utils import get_column_letter
        out = []
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                out.append(self.scalar(self.cell(sname, "{0}{1}".format(get_column_letter(c), r))))
        return out

    @staticmethod
    def scalar(v):
        """Propagate an error value stored in a cell."""
        if isinstance(v, XLError):
            raise v
        return v

    def ev(self, node, sheet):
        """Evaluate an AST node."""
        kind = node[0]
        if kind == "lit":
            return node[1]
        if kind == "ref":
            return self.ref(sheet, node[1])
        if kind == "neg":
            return broadcast(lambda a, _b: -to_num(a), self.ev(node[1], sheet), 0)
        if kind == "bin":
            op, a, b = node[1], self.ev(node[2], sheet), self.ev(node[3], sheet)
            if op == "&":
                return broadcast(lambda x, y: to_text(x) + to_text(y), a, b)
            if op in ("+", "-", "*", "/"):
                def arith(x, y):
                    x, y = to_num(x), to_num(y)
                    if op == "/" and y == 0:
                        raise XLError("#DIV/0!")
                    return {"+": x + y, "-": x - y, "*": x * y, "/": x / y if y else 0}[op]
                return broadcast(arith, a, b)
            return broadcast(lambda x, y: compare(x, y, op), a, b)
        if kind == "call":
            return self.call(node[1], node[2], sheet)
        raise ValueError("bad node {0!r}".format(node))

    def call(self, name, args, sheet):
        """Dispatch a function call (IF and IFERROR are lazy)."""
        if name == "IF":
            cond = self.ev(args[0], sheet)
            if isinstance(cond, list):
                raise XLError("#VALUE!")
            truth = bool(to_num(cond)) if not isinstance(cond, str) else bool(cond)
            if truth:
                return self.ev(args[1], sheet)
            return self.ev(args[2], sheet) if len(args) > 2 else False
        if name == "IFERROR":
            try:
                return self.ev(args[0], sheet)
            except XLError:
                return self.ev(args[1], sheet)
        vals = [self.ev(a, sheet) for a in args]
        if name == "SUM":
            return sum(float(x) for v in vals for x in flatten(v) if is_num(x))
        if name in ("MAX", "MIN"):
            nums = [float(x) for v in vals for x in flatten(v)
                    if is_num(x) or (not isinstance(v, list) and isinstance(x, bool))]
            if not nums:
                return 0.0
            return max(nums) if name == "MAX" else min(nums)
        if name == "ABS":
            return broadcast(lambda a, _b: abs(to_num(a)), vals[0], 0)
        if name in ("AND", "OR"):
            flags = [bool(to_num(x)) for v in vals for x in flatten(v) if x is not None]
            return all(flags) if name == "AND" else any(flags)
        if name == "NOT":
            return not bool(to_num(vals[0]))
        if name == "ROUNDUP":
            return roundup(to_num(vals[0]), to_num(vals[1]))
        if name == "ROUND":
            f = 10 ** int(to_num(vals[1]))
            return math.floor(to_num(vals[0]) * f + 0.5) / f
        if name == "SUMIFS":
            total = 0.0
            rng = vals[0]
            pairs = list(zip(vals[1::2], vals[2::2]))
            for i, x in enumerate(rng):
                if all(crit_match(r[i], c) for r, c in pairs) and is_num(x):
                    total += float(x)
            return total
        if name == "COUNTIFS" or name == "COUNTIF":
            pairs = list(zip(vals[0::2], vals[1::2]))
            vec = [c for _r, c in pairs if isinstance(c, list)]
            if vec:
                n = len(vec[0])
                return [float(sum(1 for i in range(len(pairs[0][0]))
                                  if all(crit_match(r[i], c[k] if isinstance(c, list) else c)
                                         for r, c in pairs)))
                        for k in range(n)]
            return float(sum(1 for i in range(len(pairs[0][0]))
                             if all(crit_match(r[i], c) for r, c in pairs)))
        if name == "COUNTBLANK":
            return float(sum(1 for x in flatten(vals[0]) if x is None or x == ""))
        if name == "SUMPRODUCT":
            arrays = [v if isinstance(v, list) else [v] for v in vals]
            n = len(arrays[0])
            total = 0.0
            for i in range(n):
                prod = 1.0
                for arr in arrays:
                    x = arr[i]
                    prod *= float(x) if is_num(x) else 0.0
                total += prod
            return total
        if name == "INDEX":
            rng = flatten(vals[0])
            idx = int(to_num(vals[1]))
            if idx < 1 or idx > len(rng):
                raise XLError("#REF!")
            return rng[idx - 1]
        if name == "MATCH":
            for i, x in enumerate(flatten(vals[1]), start=1):
                if x is not None and compare(x, vals[0], "="):
                    return float(i)
            raise XLError("#N/A")
        if name == "WORKDAY":
            return workday(to_num(vals[0]), to_num(vals[1]))
        if name == "ISNUMBER":
            return is_num(vals[0])
        if name == "LEFT":
            return to_text(vals[0])[: int(to_num(vals[1]))]
        raise ValueError("function not supported by the evaluator: {0}".format(name))


def serial_to_date(serial):
    """Convert an Excel serial number to a date."""
    return (EXCEL_EPOCH + _dt.timedelta(days=int(serial))).date()


def date_to_serial(d):
    """Convert a date to an Excel serial number."""
    return float((_dt.datetime(d.year, d.month, d.day) - EXCEL_EPOCH).days)
