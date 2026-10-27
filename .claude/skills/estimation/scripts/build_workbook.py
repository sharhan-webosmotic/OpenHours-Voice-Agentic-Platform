#!/usr/bin/env python3
"""build_workbook.py - render the simple estimation workbook (schemaVersion 3, v6.1 releases).

Writes exactly five sheets, described in EXCEL-FORMAT.md:

    Summary | Features | Tasks | Recommended | Validation

Tasks is the only sheet with typed hours. Every other number is a live Excel
formula (SUMIFS, SUM, COUNTIFS, SUMPRODUCT, WORKDAY, IF), so editing a task
updates every total. There is no margin, no markup and no buffer anywhere.

Usage:
    python3 build_workbook.py <input.json> <output.xlsx>
    python3 build_workbook.py --selftest
"""

from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import math
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from formula_eval import Evaluator, date_to_serial, workday

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
except ImportError:  # pragma: no cover - environment guard
    print("ERROR: openpyxl is required. Install it with: python3 -m pip install openpyxl",
          file=sys.stderr)
    raise

BLUE = "4472C4"
DARK = "1F3864"
GREY = "A6A6A6"
LIGHT = "F2F2F2"
YELLOW = "FFF2A8"
PASS_FILL, FAIL_FILL, REVIEW_FILL = "C6EFCE", "FFC7CE", "FFEB9C"

WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
INPUT_FONT = Font(bold=True, color="0000CC")
MUTED = Font(color="808080", italic=True)
HOURS_FMT = "0.0"
DATE_FMT = "yyyy-mm-dd"

FIXED_ROLES = [("UIUX", "UI/UX"), ("Frontend", "Frontend"), ("Backend", "Backend"),
               ("AI", "AI"), ("Mobile", "Mobile"), ("DevOps", "DevOps"), ("QA", "QA"),
               ("PM", "PM")]
PRIORITIES = ["P0", "P1", "P2", "P3"]
RECOMMENDATIONS = ["Build Now", "Build Next", "Defer", "Skip", "Needs Clarification"]
SCOPES = ["Included", "Deferred", "Excluded", "Pending"]
SOURCES = ["Client Requirement", "Recommended", "Future Enhancement"]
ESTIMATE_STATUSES = ["Confirmed", "Provisional", "Awaiting Clarification"]
CONFIDENCES = ["High", "Medium", "Low"]
COMMITMENTS = ["Committed", "Tentative"]
DEFAULT_RELEASE = "Full scope"
DEFAULT_COMMITMENT = "Committed"
SCOPE_CASES = {
    "A": "A: Full quote (one release, everything committed)",
    "B": "B: Client-defined MVP (MVP committed, later phases tentative)",
    "C": "C: MVP-first, the default (proposed MVP committed, later phases tentative)",
}
PENDING = "Pending scheduling inputs"
ETA_LABEL = ("Provisional: roles in parallel within a milestone, milestones in sequence, "
             "committed milestones first and tentative milestones after")
ETA_METHOD = "Per-role: working days = max over roles of role hours / (people x hours per day)"
EPS = 0.001

FEATURE_FIXED_COLS = ["ID", "Name", "Description", "Source", "Priority", "Recommendation",
                      "Reasoning", "Milestone", "Scope Status", "Release", "Commitment",
                      "Scope Decision"]
TASK_COLS = ["Feature ID", "Milestone", "Task", "Role", "Min Hours", "Max Hours",
             "Confidence", "Assumption", "Scope Status (formula)", "Release (formula)",
             "Commitment (formula)"]
REC_COLS = ["ID", "Feature", "Category", "Problem", "Priority", "Release", "Effort Hours",
            "Depends On", "Risk If Omitted", "Rationale", "Approval", "Included in Estimate"]
CHECKS = [
    "UI/UX hours reconcile across Tasks, Features and Summary",
    "Feature role hours reconcile with Tasks",
    "Feature total equals the sum of its role columns",
    "Milestone totals equal project totals (each role)",
    "ETA is not inferred from total effort alone",
    "Every included feature has a priority",
    "Excluded and deferred features are absent from committed hours",
    "Recommended items are not in committed hours",
    "No task counted twice (Feature + Task + Role)",
    "No margin or markup is applied",
    "Unknown estimates are not treated as zero",
    "Every task role is a known role and every task feature exists",
    "Tentative hours are absent from the committed total",
    "Every committed feature is P0 or P1, or its reasoning explains why",
    "Every milestone has a release and a commitment",
]


def fill(hex_color: str) -> PatternFill:
    """Return a solid fill for a hex colour."""
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def header(ws: Any, row: int, labels: list[str], color: str = BLUE, start: int = 1) -> None:
    """Write a styled header row."""
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=label)
        c.font = WHITE_BOLD
        c.fill = fill(color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def section(ws: Any, row: int, text: str) -> None:
    """Write a bold section title."""
    ws.cell(row=row, column=1, value=text).font = Font(bold=True, size=12, color=DARK)


def as_input(cell: Any) -> None:
    """Style a cell as a user-editable input (blue text on yellow)."""
    cell.font = INPUT_FONT
    cell.fill = fill(YELLOW)


def esc(text: str) -> str:
    """Escape a string literal for use inside an Excel formula."""
    return '"' + str(text).replace('"', '""') + '"'


def deps_list(value: Any) -> list[str]:
    """Normalise dependsOn (list or comma string) to a list of IDs."""
    if not value:
        return []
    if isinstance(value, str):
        value = value.split(",")
    return [str(v).strip() for v in value if str(v).strip()]


def roles_of(data: dict) -> list[tuple[str, str]]:
    """Fixed roles in contract order, then each otherRoles entry."""
    extra = [(r["id"], r.get("name") or r["id"]) for r in data.get("otherRoles", [])]
    return FIXED_ROLES + extra


def ms_release(m: dict) -> str:
    """Effective release label of a milestone (defaults to Full scope)."""
    return m.get("release") or DEFAULT_RELEASE


def ms_commitment(m: dict) -> str:
    """Effective commitment of a milestone (defaults to Committed)."""
    return m.get("commitment") or DEFAULT_COMMITMENT


def column_order(milestones: list[dict]) -> list[dict]:
    """Committed milestones first, then tentative, grouped by release, JSON order inside."""
    first_seen: dict[str, int] = {}
    for i, m in enumerate(milestones):
        first_seen.setdefault(ms_release(m), i)
    ranked = sorted(enumerate(milestones), key=lambda p: (
        0 if ms_commitment(p[1]) == "Committed" else 1, first_seen[ms_release(p[1])], p[0]))
    return [m for _i, m in ranked]


def tentative_releases(milestones: list[dict]) -> list[str]:
    """Release labels of tentative milestones, in column order, without duplicates."""
    out: list[str] = []
    for m in column_order(milestones):
        if ms_commitment(m) == "Tentative" and ms_release(m) not in out:
            out.append(ms_release(m))
    return out


def check_input(data: dict) -> tuple[list[str], list[str]]:
    """Return (errors, warnings). Errors stop the build, warnings surface on Validation."""
    errors, warnings = [], []
    if data.get("schemaVersion") != 3:
        errors.append("schemaVersion must be 3")
    for key in ("project", "milestones", "features", "tasks"):
        if key not in data:
            errors.append("missing top-level key: " + key)
    if errors:
        return errors, warnings
    if not data["milestones"]:
        errors.append("at least one milestone is required")
    if not data["tasks"]:
        errors.append("at least one task is required")
    enums = [("source", SOURCES), ("priority", PRIORITIES), ("recommendation", RECOMMENDATIONS),
             ("scopeStatus", SCOPES), ("estimateStatus", ESTIMATE_STATUSES)]
    ms_ids = {m["id"] for m in data["milestones"]}
    for m in data["milestones"]:
        if m.get("commitment") not in (None, "", *COMMITMENTS):
            warnings.append("{0}: commitment {1!r} not in {2}".format(m["id"], m.get("commitment"), COMMITMENTS))
        if not m.get("release") or not m.get("commitment"):
            warnings.append("{0}: release or commitment missing, defaulted to {1} / {2}".format(
                m["id"], ms_release(m), ms_commitment(m)))
    case = data["project"].get("scopeCase")
    if case is not None and case not in SCOPE_CASES:
        warnings.append("project.scopeCase {0!r} not in A, B, C".format(case))
    for f in data["features"]:
        if f.get("scopeDecision") not in (None, "", *COMMITMENTS):
            warnings.append("{0}: scopeDecision {1!r} not in {2}".format(f.get("id"), f.get("scopeDecision"), COMMITMENTS))
        for key, allowed in enums:
            if f.get(key) not in allowed:
                warnings.append("{0}: {1}={2!r} not in {3}".format(f.get("id"), key, f.get(key), allowed))
        if f.get("milestone") not in ms_ids:
            warnings.append("{0}: milestone {1!r} is not defined".format(f.get("id"), f.get("milestone")))
        if f.get("recommendation") == "Skip" and f.get("scopeStatus") == "Included":
            warnings.append("{0}: recommendation Skip but scopeStatus Included".format(f.get("id")))
    role_ids = {r for r, _ in roles_of(data)}
    feat_ids = {f["id"] for f in data["features"]}
    for i, t in enumerate(data["tasks"], start=1):
        if t.get("role") not in role_ids:
            warnings.append("task {0}: unknown role {1!r}".format(i, t.get("role")))
        if t.get("feature") != "SHARED" and t.get("feature") not in feat_ids:
            warnings.append("task {0}: unknown feature {1!r}".format(i, t.get("feature")))
        if t.get("feature") == "SHARED" and t.get("milestone") not in ms_ids:
            warnings.append("task {0}: SHARED task needs a defined milestone".format(i))
        if t.get("confidence") not in CONFIDENCES:
            warnings.append("task {0}: confidence {1!r} not in {2}".format(i, t.get("confidence"), CONFIDENCES))
        mn, mx = t.get("minHours"), t.get("maxHours")
        if mn is not None and mx is not None and mx < mn:
            warnings.append("task {0}: maxHours below minHours".format(i))
    return errors, warnings


class Layout:
    """Cell addresses shared between sheet builders, the verifier and the CLI."""

    def __init__(self, data: dict) -> None:
        self.roles = roles_of(data)
        self.ms = data["milestones"]
        self.order = column_order(self.ms)
        self.ms_col = {m["id"]: 3 + j for j, m in enumerate(self.order)}
        self.tent_rels = tentative_releases(self.ms)
        self.m = max(len(data["features"]), 1)
        self.n = len(data["tasks"])
        self.k = len(data.get("recommended", []))
        self.f_last = self.m + 1
        self.t_last = self.n + 1
        self.r_last = max(self.k, 1) + 1
        first_role = len(FEATURE_FIXED_COLS) + 1
        self.role_col = {key: first_role + i for i, (key, _) in enumerate(self.roles)}
        self.f_total = first_role + len(self.roles)
        self.f_pd, self.f_deps, self.f_est, self.f_eta = (self.f_total + 1, self.f_total + 2,
                                                          self.f_total + 3, self.f_total + 4)
        self.cells: dict[str, str] = {}

    def trng(self, col: str) -> str:
        """Absolute Tasks range for one column letter."""
        return "Tasks!${0}$2:${0}${1}".format(col, self.t_last)

    def frng(self, col: int | str) -> str:
        """Absolute Features range for one column (index or letter)."""
        col = L(col) if isinstance(col, int) else col
        return "Features!${0}$2:${0}${1}".format(col, self.f_last)

    def lookup(self, key_ref: str, field: str) -> str:
        """Formula body: a milestone's effective release or commitment, by milestone ID."""
        return 'IFERROR(INDEX({0},MATCH({1},{2},0)),"Unknown milestone")'.format(
            self.cells[field], key_ref, self.cells["ms_id"])


def build_tasks(wb: Any, data: dict, lay: Layout) -> None:
    """Tasks sheet: the only typed hours, plus a scope-status lookup helper."""
    ws = wb.create_sheet("Tasks")
    header(ws, 1, TASK_COLS[:8])
    header(ws, 1, TASK_COLS[8:], GREY, start=9)
    for r, t in enumerate(data["tasks"], start=2):
        ws.cell(row=r, column=1, value=t.get("feature"))
        if t.get("feature") == "SHARED":
            ws.cell(row=r, column=2, value=t.get("milestone"))
        else:
            ws.cell(row=r, column=2, value='=IFERROR(INDEX({0},MATCH(A{1},{2},0)),"Unknown feature")'
                    .format(lay.frng("H"), r, lay.frng("A")))
        ws.cell(row=r, column=3, value=t.get("task"))
        ws.cell(row=r, column=4, value=t.get("role"))
        for col, key in ((5, "minHours"), (6, "maxHours")):
            c = ws.cell(row=r, column=col, value=t.get(key))
            c.number_format = HOURS_FMT
            as_input(c)
        ws.cell(row=r, column=7, value=t.get("confidence"))
        ws.cell(row=r, column=8, value=t.get("assumption"))
        ws.cell(row=r, column=9, value='=IF(A{0}="SHARED","Included",IFERROR(INDEX({1},MATCH(A{0},{2},0)),'
                '"Unknown feature"))'.format(r, lay.frng("I"), lay.frng("A"))).font = MUTED
        ws.cell(row=r, column=10, value="=" + lay.lookup("B{0}".format(r), "ms_rel")).font = MUTED
        ws.cell(row=r, column=11, value="=" + lay.lookup("B{0}".format(r), "ms_com")).font = MUTED
    for col, w in zip("ABCDEFGHIJK", (11, 10, 42, 11, 10, 10, 11, 48, 16, 14, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def build_features(wb: Any, data: dict, lay: Layout) -> None:
    """Features sheet: one row per feature, every hour cell a SUMIFS over Tasks."""
    ws = wb.create_sheet("Features")
    labels = FEATURE_FIXED_COLS + [lbl for _, lbl in lay.roles] + [
        "Total Hours", "Person-Days", "Depends On", "Estimate Status", "Calendar ETA"]
    header(ws, 1, labels)
    hpd = lay.cells["hpd"]
    for r, f in enumerate(data["features"], start=2):
        vals = [f.get("id"), f.get("name"), f.get("description"), f.get("source"), f.get("priority"),
                f.get("recommendation"), f.get("reasoning"), f.get("milestone"), f.get("scopeStatus")]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
        ws.cell(row=r, column=10, value="=" + lay.lookup("$H{0}".format(r), "ms_rel"))
        ws.cell(row=r, column=11, value="=" + lay.lookup("$H{0}".format(r), "ms_com"))
        ws.cell(row=r, column=12, value=f.get("scopeDecision") or None)
        for key, col in lay.role_col.items():
            ws.cell(row=r, column=col, value="=SUMIFS({0},{1},$A{2},{3},{4})".format(
                lay.trng("E"), lay.trng("A"), r, lay.trng("D"), esc(key))).number_format = HOURS_FMT
        first, last = L(lay.role_col[lay.roles[0][0]]), L(lay.role_col[lay.roles[-1][0]])
        tc = ws.cell(row=r, column=lay.f_total, value="=SUM({0}{2}:{1}{2})".format(first, last, r))
        tc.number_format, tc.font = HOURS_FMT, BOLD
        ws.cell(row=r, column=lay.f_pd, value="={0}{1}/{2}".format(L(lay.f_total), r, hpd)).number_format = "0.00"
        ws.cell(row=r, column=lay.f_deps, value=", ".join(deps_list(f.get("dependsOn"))))
        ws.cell(row=r, column=lay.f_est, value=f.get("estimateStatus"))
        ws.cell(row=r, column=lay.f_eta, value=PENDING).font = MUTED
    widths = [7, 26, 40, 18, 9, 18, 40, 10, 12, 14, 12, 14] + [9] * len(lay.roles) + [11, 11, 14, 20, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = "C2"


def build_recommended(wb: Any, data: dict, lay: Layout) -> None:
    """Recommended register: never part of committed hours."""
    ws = wb.create_sheet("Recommended")
    header(ws, 1, REC_COLS)
    keys = ["id", "feature", "category", "problem", "priority", "release", "effortHours",
            "dependsOn", "riskIfOmitted", "rationale", "approval"]
    for r, item in enumerate(data.get("recommended", []), start=2):
        for c, key in enumerate(keys, start=1):
            v = item.get(key)
            ws.cell(row=r, column=c, value=", ".join(v) if isinstance(v, list) else v)
        ws.cell(row=r, column=7).number_format = HOURS_FMT
        ws.cell(row=r, column=12, value="No")
    note_row = lay.r_last + 2
    ws.cell(row=note_row, column=1, value=(
        "Recommended items never add to committed hours. To include one, add it as a feature "
        "(source Recommended) with real tasks on the Tasks sheet.")).font = MUTED
    for i, w in enumerate((7, 26, 14, 40, 9, 14, 11, 14, 12, 44, 11, 12), start=1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = "A2"


def build_milestone_table(ws: Any, data: dict, lay: Layout, r: int) -> int:
    """Releases and milestones table (typed release and commitment, effective values). Returns next free row."""
    section(ws, r, "Releases and milestones (delivery order)")
    header(ws, r + 1, ["Milestone", "Name", "Release (input)", "Commitment (input)", "Release", "Commitment"])
    first = r + 2
    for i, m in enumerate(lay.ms):
        row = first + i
        ws.cell(row=row, column=1, value=m["id"])
        ws.cell(row=row, column=2, value=m.get("name"))
        for col, key in ((3, "release"), (4, "commitment")):
            as_input(ws.cell(row=row, column=col, value=m.get(key) or None))
        ws.cell(row=row, column=5, value='=IF(C{0}="",{1},C{0})'.format(row, esc(DEFAULT_RELEASE)))
        ws.cell(row=row, column=6, value='=IF(D{0}="",{1},D{0})'.format(row, esc(DEFAULT_COMMITMENT))).font = BOLD
        lay.cells["ms_row_" + m["id"]] = str(row)
    last = first + len(lay.ms) - 1
    ws.cell(row=last + 1, column=1, value=(
        "A blank Release defaults to {0} and a blank Commitment to {1}, which turns Validation check 15 "
        "to REVIEW. Tentative milestones are future phases: estimated, shown, never added to the committed "
        "number.".format(DEFAULT_RELEASE, DEFAULT_COMMITMENT))).font = MUTED
    for key, col in (("ms_id", "A"), ("ms_rel_in", "C"), ("ms_com_in", "D"), ("ms_rel", "E"), ("ms_com", "F")):
        lay.cells[key] = "Summary!${0}${1}:${0}${2}".format(col, first, last)
    return last + 3


def build_role_table(ws: Any, data: dict, lay: Layout, r: int) -> int:
    """Role x milestone hours, milestone columns grouped by release, separate subtotals. Returns next free row."""
    section(ws, r, "Effort by role and milestone (hours): committed and tentative are separate")
    n_ms = len(lay.order)
    c_col, t_col, a_col, pd_col = 3 + n_ms, 4 + n_ms, 5 + n_ms, 6 + n_ms
    rel_row, com_row, hdr = r + 1, r + 2, r + 3
    for label, row, field in (("Release", rel_row, "E"), ("Commitment", com_row, "F")):
        ws.cell(row=row, column=1, value=label).font = BOLD
        for m in lay.order:
            c = ws.cell(row=row, column=lay.ms_col[m["id"]],
                        value="=${0}${1}".format(field, lay.cells["ms_row_" + m["id"]]))
            c.font, c.alignment = BOLD, Alignment(horizontal="center")
    header(ws, hdr, ["Role", "Key"] + ["{0}: {1}".format(m["id"], m["name"]) for m in lay.order]
           + ["Committed subtotal", "Tentative subtotal", "Committed + tentative (not a commitment)",
              "Committed Person-Days"])
    for col in (t_col, a_col):
        ws.cell(row=hdr, column=col).fill = fill(GREY)
    t_e, t_d, t_i, t_k = lay.trng("E"), lay.trng("D"), lay.trng("I"), lay.trng("K")
    r0 = hdr + 1
    for i, (key, label) in enumerate(lay.roles):
        row = r0 + i
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=key).font = MUTED
        for m in lay.order:
            ws.cell(row=row, column=lay.ms_col[m["id"]], value='=SUMIFS({0},{1},{2},{3},$B{4},{5},"Included")'.format(
                t_e, lay.trng("B"), esc(m["id"]), t_d, row, t_i)).number_format = HOURS_FMT
        for col, crit in ((c_col, ',{0},"Committed"'.format(t_k)), (t_col, ',{0},"Tentative"'.format(t_k)),
                          (a_col, "")):
            ws.cell(row=row, column=col, value='=SUMIFS({0},{1},$B{2},{3},"Included"{4})'.format(
                t_e, t_d, row, t_i, crit)).number_format = HOURS_FMT
        ws.cell(row=row, column=pd_col, value="={0}{1}/$B$5".format(L(c_col), row)).number_format = "0.00"
    r_end = r0 + len(lay.roles) - 1
    tot = r_end + 1
    ws.cell(row=tot, column=1, value="Total").font = BOLD
    for col in list(lay.ms_col.values()) + [c_col, t_col, a_col]:
        c = ws.cell(row=tot, column=col, value="=SUM({0}{1}:{0}{2})".format(L(col), r0, r_end))
        c.number_format, c.font = HOURS_FMT, BOLD
    ws.cell(row=tot, column=pd_col, value="={0}{1}/$B$5".format(L(c_col), tot)).number_format = "0.00"
    lay.cells.update(
        role_first=str(r0), role_last=str(r_end), total_row=str(tot), rel_row=str(rel_row), com_row=str(com_row),
        ms_first=L(3), ms_last=L(2 + n_ms), c_col=L(c_col), t_col=L(t_col), proj_col=L(a_col),
        approved="Summary!${0}${1}".format(L(c_col), tot), tent_total="Summary!${0}${1}".format(L(t_col), tot),
        all_incl="Summary!${0}${1}".format(L(a_col), tot), uiux_summary="Summary!${0}${1}".format(L(a_col), r0))
    return tot + 2


def build_buckets(ws: Any, data: dict, lay: Layout, r: int) -> int:
    """Scope buckets: committed, each tentative release, proposed, deferred, excluded, unestimated."""
    section(ws, r, "Scope buckets (each bucket is separate, only Committed is a commitment)")
    header(ws, r + 1, ["Bucket", "Hours", "Person-Days", "Upper Range", "Note"])
    t_e, t_f, t_i, t_j, t_k = lay.trng("E"), lay.trng("F"), lay.trng("I"), lay.trng("J"), lay.trng("K")

    def upper(extra: str) -> str:
        return '=SUMIFS({0},{1},"Included"{3})+SUMIFS({2},{1},"Included"{3},{0},"")'.format(t_f, t_i, t_e, extra)

    committed_rels = []
    for m in lay.order:
        if ms_commitment(m) == "Committed" and ms_release(m) not in committed_rels:
            committed_rels.append(ms_release(m))
    com = ',{0},"Committed"'.format(t_k)
    tent = ',{0},"Tentative"'.format(t_k)
    rec_sum = "=SUM(Recommended!$G$2:$G${0})".format(lay.r_last) if lay.k else "=0"
    buckets = [("approved", "Committed ({0})".format(", ".join(committed_rels) or "none"),
                "=" + lay.cells["approved"].replace("Summary!", ""), upper(com),
                "Included features + SHARED work in Committed milestones")]
    for i, rel in enumerate(lay.tent_rels):
        crit = ',{0},{1},{2},"Tentative"'.format(t_j, esc(rel), t_k)
        buckets.append(("tent_{0}".format(i), "Future phase (tentative): {0}".format(rel),
                        '=SUMIFS({0},{1},"Included"{2})'.format(t_e, t_i, crit), upper(crit),
                        "Estimated like the MVP, not a commitment"))
    buckets += [
        ("tent_total", "Future phases total (tentative, not a commitment)",
         "=" + lay.cells["tent_total"].replace("Summary!", ""), upper(tent), "Never added to the committed number"),
        ("prop_feat", "Proposed: Pending features", '=SUMIFS({0},{1},"Pending")'.format(t_e, t_i), None,
         "Awaiting a scope decision"),
        ("prop_rec", "Proposed: recommended items with effort", rec_sum, None, "Register only, never committed"),
        ("prop_total", "Proposed total", None, None, "Not committed"),
        ("deferred", "Deferred", '=SUMIFS({0},{1},"Deferred")'.format(t_e, t_i), None, "Not planned in any phase"),
        ("excluded", "Excluded", '=SUMIFS({0},{1},"Excluded")'.format(t_e, t_i), None, "Out of scope"),
        ("unest_tasks", "Unestimated tasks (count, not hours)", "=COUNTBLANK({0})".format(t_e), None,
         "Tasks with no Min Hours"),
        ("unest_feats", "Features with no tasks (count, not hours)",
         "=SUMPRODUCT(--(COUNTIF({0},{1})=0))".format(lay.trng("A"), lay.frng("A")), None, "Not yet decomposed"),
        ("upper", "Upper range (not committed)", upper(com), None,
         "Max Hours where given, else Min Hours, committed scope only"),
    ]
    rr = r + 2
    for key, label, formula, up, note in buckets:
        ws.cell(row=rr, column=1, value=label).font = BOLD if key in ("approved", "tent_total") else Font()
        if key == "prop_total":
            formula = "=B{0}+B{1}".format(lay.cells["b_prop_feat"].split("$")[-1], lay.cells["b_prop_rec"].split("$")[-1])
        c = ws.cell(row=rr, column=2, value=formula)
        if key.startswith("unest"):
            c.number_format = "0"
            ws.cell(row=rr, column=3, value="count")
        else:
            c.number_format = HOURS_FMT
            ws.cell(row=rr, column=3, value="=B{0}/$B$5".format(rr)).number_format = "0.00"
        if up:
            ws.cell(row=rr, column=4, value=up).number_format = HOURS_FMT
            lay.cells["u_" + key] = "Summary!$D${0}".format(rr)
        ws.cell(row=rr, column=5, value=note).font = MUTED
        lay.cells["b_" + key] = "Summary!$B${0}".format(rr)
        rr += 1
    return rr + 1


def build_summary(wb: Any, data: dict, lay: Layout) -> None:
    """Summary: header, releases, role x milestone table, buckets, counts, dependencies, ETA."""
    ws = wb.active
    ws.title = "Summary"
    p = data["project"]
    ws["A1"] = "{0}: effort estimate".format(p.get("name", "Project"))
    ws["A1"].font = TITLE
    for r, (label, val) in enumerate((("Client", p.get("client")), ("Prepared by", p.get("preparedBy")),
                                      ("Date", p.get("date"))), start=2):
        ws.cell(row=r, column=1, value=label).font = BOLD
        ws.cell(row=r, column=2, value=val)
    ws["A5"], ws["A5"].font = "Hours per day", BOLD
    ws["B5"] = p.get("hoursPerDay", 8)
    as_input(ws["B5"])
    lay.cells["hpd"] = "Summary!$B$5"
    ws["A6"] = ("Committed hours are the minimum realistic hours per task with no margin, no markup and no "
                "buffer. Committed = Included features + SHARED work in Committed milestones. Future phases are "
                "tentative, shown with their own hours, and never added to the committed number.")
    ws["A6"].font = MUTED
    ws["A7"], ws["A7"].font = "Scope case", BOLD
    ws["B7"] = SCOPE_CASES.get(p.get("scopeCase"), "Not recorded")

    r = build_milestone_table(ws, data, lay, 9)
    r = build_role_table(ws, data, lay, r)
    rr = build_buckets(ws, data, lay, r)

    r = rr
    section(ws, r, "Feature counts")
    header(ws, r + 1, ["Priority", "Features", "", "Recommendation", "Features"])
    rows = max(len(PRIORITIES), len(RECOMMENDATIONS))
    for i in range(rows):
        row = r + 2 + i
        if i < len(PRIORITIES):
            ws.cell(row=row, column=1, value=PRIORITIES[i])
            ws.cell(row=row, column=2, value='=COUNTIF({0},"{1}")'.format(lay.frng("E"), PRIORITIES[i]))
            lay.cells["cnt_" + PRIORITIES[i]] = "Summary!$B${0}".format(row)
        ws.cell(row=row, column=4, value=RECOMMENDATIONS[i])
        ws.cell(row=row, column=5, value='=COUNTIF({0},"{1}")'.format(lay.frng("F"), RECOMMENDATIONS[i]))
        lay.cells["cnt_" + RECOMMENDATIONS[i]] = "Summary!$E${0}".format(row)

    r = r + 2 + rows + 1
    section(ws, r, "Feature dependencies (unresolved when a dependency is not Included or does not exist)")
    with_deps = [f for f in data["features"] if deps_list(f.get("dependsOn"))]
    header(ws, r + 1, ["Feature", "Depends On", "Status"])
    rr = r + 2
    dep_first = rr
    for f in with_deps:
        conds = ['IFERROR(INDEX({0},MATCH({1},{2},0)),"Missing")="Included"'.format(
            lay.frng("I"), esc(d), lay.frng("A")) for d in deps_list(f["dependsOn"])]
        ws.cell(row=rr, column=1, value=f["id"])
        ws.cell(row=rr, column=2, value=", ".join(deps_list(f["dependsOn"])))
        ws.cell(row=rr, column=3, value='=IF(AND({0}),"Resolved","Unresolved")'.format(",".join(conds)))
        lay.cells["dep_" + f["id"]] = "Summary!$C${0}".format(rr)
        rr += 1
    if not with_deps:
        ws.cell(row=rr, column=1, value="No feature dependencies recorded").font = MUTED
        rr += 1
    ws.cell(row=rr, column=1, value="Features with an unresolved dependency").font = BOLD
    ws.cell(row=rr, column=2, value='=COUNTIF($C${0}:$C${1},"Unresolved")'.format(dep_first, max(rr - 1, dep_first))
            if with_deps else 0)
    lay.cells["dep_unresolved"] = "Summary!$B${0}".format(rr)

    build_eta(ws, data, lay, rr + 2)
    for col, w in zip("ABCDEFGHIJ", (44, 14, 16, 22, 14, 14, 14, 14, 14, 14)):
        ws.column_dimensions[col].width = w


def _eta_delivery_rows(ws: Any, lay: Layout, first: int, schedule: bool) -> int:
    """Committed delivery date and tentative phases end rows under the ETA table. Returns the first of them."""
    order = lay.order
    com_rows = [first + i for i, m in enumerate(order) if ms_commitment(m) == "Committed"]
    tent_rows = [first + i for i, m in enumerate(order) if ms_commitment(m) != "Committed"]
    row = first + len(order)
    rows = (("Committed delivery date (provisional)", com_rows, "None committed", "eta_delivery"),
            ("Tentative phases end (provisional, not a commitment)", tent_rows, "No tentative phases",
             "eta_tent_end"))
    for i, (label, rows_of, empty, key) in enumerate(rows):
        ws.cell(row=row + i, column=1, value=label).font = BOLD
        if not schedule:
            ws.cell(row=row + i, column=2, value=PENDING)
            lay.cells[key] = "Summary!$B${0}".format(row + i)
            continue
        c = ws.cell(row=row + i, column=4, value="=D{0}".format(rows_of[-1]) if rows_of else empty)
        c.number_format = DATE_FMT
        lay.cells[key] = "Summary!$D${0}".format(row + i)
    return row


def build_eta(ws: Any, data: dict, lay: Layout, r: int) -> None:
    """Provisional milestone ETA table (committed first, tentative after), or Pending scheduling inputs."""
    section(ws, r, "Provisional milestone ETA")
    ws.cell(row=r + 1, column=1, value=ETA_LABEL).font = MUTED
    ws.cell(row=r + 2, column=1, value="ETA method").font = BOLD
    sched = data.get("schedule")
    lay.cells["eta_method"] = "Summary!$B${0}".format(r + 2)
    cols = ["Milestone", "Working Days", "Start", "End", "Staffing Gaps", "Commitment"]
    if not sched:
        ws.cell(row=r + 2, column=2, value=PENDING)
        header(ws, r + 4, cols)
        first = r + 5
        for i, m in enumerate(lay.order):
            ws.cell(row=first + i, column=1, value="{0}: {1}".format(m["id"], m["name"]))
            for col in (2, 3, 4):
                ws.cell(row=first + i, column=col, value=PENDING).font = MUTED
            ws.cell(row=first + i, column=6, value="=$F${0}".format(lay.cells["ms_row_" + m["id"]]))
        last = first + len(lay.order) - 1
        _eta_delivery_rows(ws, lay, first, False)
        lay.cells["eta_range"] = "Summary!$B${0}:$D${1}".format(first, last)
        return
    ws.cell(row=r + 2, column=2, value=ETA_METHOD)
    ws.cell(row=r + 3, column=1, value="Start date").font = BOLD
    sd = ws.cell(row=r + 3, column=2, value=dt.date.fromisoformat(sched["startDate"]))
    sd.number_format = DATE_FMT
    as_input(sd)
    start_ref = "$B${0}".format(r + 3)
    header(ws, r + 5, ["Team (people per role)", "Key", "People"])
    people = sched.get("people", {})
    p_first = r + 6
    for i, (key, label) in enumerate(lay.roles):
        ws.cell(row=p_first + i, column=1, value=label)
        ws.cell(row=p_first + i, column=2, value=key).font = MUTED
        c = ws.cell(row=p_first + i, column=3, value=people.get(key, 0))
        as_input(c)
    p_last = p_first + len(lay.roles) - 1
    m_hdr = p_last + 2
    header(ws, m_hdr, cols)
    rf, rl = int(lay.cells["role_first"]), int(lay.cells["role_last"])
    first = m_hdr + 1
    for i, m in enumerate(lay.order):
        row = first + i
        mcol = L(lay.ms_col[m["id"]])
        parts = ["IF($C${0}>0,{1}{2}/($C${0}*$B$5),0)".format(p_first + j, mcol, rf + j)
                 for j in range(len(lay.roles))]
        ws.cell(row=row, column=1, value="{0}: {1}".format(m["id"], m["name"]))
        ws.cell(row=row, column=2, value="=ROUNDUP(MAX(0,{0}),0)".format(",".join(parts)))
        start = "=WORKDAY({0}-1,1)".format(start_ref) if i == 0 else "=WORKDAY(D{0},1)".format(row - 1)
        ws.cell(row=row, column=3, value=start).number_format = DATE_FMT
        ws.cell(row=row, column=4, value="=IF(B{0}=0,C{0},WORKDAY(C{0},B{0}-1))".format(row)).number_format = DATE_FMT
        ws.cell(row=row, column=5, value="=SUMPRODUCT(({0}{1}:{0}{2}>0)*($C${3}:$C${4}=0))".format(
            mcol, rf, rl, p_first, p_last))
        ws.cell(row=row, column=6, value="=$F${0}".format(lay.cells["ms_row_" + m["id"]]))
        lay.cells["eta_end_" + m["id"]] = "Summary!$D${0}".format(row)
        lay.cells["eta_days_" + m["id"]] = "Summary!$B${0}".format(row)
    last = first + len(lay.order) - 1
    d_row = _eta_delivery_rows(ws, lay, first, True)
    ws.cell(row=d_row, column=5, value="=SUM(E{0}:E{1})".format(first, last)).font = BOLD
    lay.cells["eta_gaps"] = "Summary!$E${0}".format(d_row)


def release_checks(lay: Layout) -> list[tuple[str, str]]:
    """Checks 13 to 15: committed vs tentative separation, committed priorities, milestone release fields."""
    c, t = lay.cells, lay.trng
    approved, tent_total = c["approved"], c["tent_total"]
    tot = c["total_row"]
    ms_row = "Summary!${0}${1}:${2}${1}".format(c["ms_first"], tot, c["ms_last"])
    com_row = "Summary!${0}${1}:${2}${1}".format(c["ms_first"], c["com_row"], c["ms_last"])

    def independent(word: str) -> str:
        return 'SUMPRODUCT(({0}="Included")*(COUNTIFS({1},{2},{3},"{4}")>0)*{5})'.format(
            t("I"), c["ms_id"], t("B"), c["ms_com"], word, t("E"))

    leak = 'COUNTIFS({0},"Included",{1},"Committed",{2},"Tentative")'.format(
        lay.frng("I"), lay.frng("K"), lay.frng("L"))
    ind_c, ind_t = independent("Committed"), independent("Tentative")
    check13 = ('=IF(AND(ABS({0}-{1})<{5},ABS({2}-{3})<{5},ABS({0}-SUMIFS({6},{7},"Committed"))<{5},{4}=0),'
               '"PASS","FAIL")').format(approved, ind_c, tent_total, ind_t, leak, EPS, ms_row, com_row)
    detail13 = ('="Committed "&{0}&" h, tentative "&{1}&" h kept separate | "&{2}&" feature(s) planned as '
                'Tentative sit in a Committed milestone"').format(approved, tent_total, leak)
    low = ('SUMPRODUCT(({0}="Included")*({1}="Committed")*({2}<>"P0")*({2}<>"P1")*({3}=""))').format(
        lay.frng("I"), lay.frng("K"), lay.frng("E"), lay.frng("G"))
    check14 = '=IF({0}=0,"PASS","REVIEW")'.format(low)
    detail14 = '={0}&" committed feature(s) are P2 or P3 (or unprioritised) with no reasoning"'.format(low)
    n = len(lay.ms)
    blanks = "(COUNTBLANK({0})+COUNTBLANK({1}))".format(c["ms_rel_in"], c["ms_com_in"])
    valid = '(COUNTIF({0},"Committed")+COUNTIF({0},"Tentative"))'.format(c["ms_com"])
    check15 = '=IF(AND({0}=0,{1}={2}),"PASS","REVIEW")'.format(blanks, valid, n)
    detail15 = ('={0}&" release or commitment cell(s) defaulted, "&({2}-{1})&" commitment value(s) are '
                'neither Committed nor Tentative"').format(blanks, valid, n)
    return [(check13, detail13), (check14, detail14), (check15, detail15)]


def build_validation(wb: Any, data: dict, lay: Layout) -> None:
    """Validation sheet: fifteen live checks, PASS / FAIL / REVIEW."""
    ws = wb.create_sheet("Validation")
    ws["A1"] = "Validation (live formulas: a math mismatch is FAIL, a missing input is REVIEW)"
    ws["A1"].font = TITLE
    header(ws, 3, ["#", "Check", "Result", "Detail"])
    c = lay.cells
    t = lay.trng
    approved, all_incl = c["approved"], c["all_incl"]
    rf, tot, pc = c["role_first"], c["total_row"], c["proj_col"]
    role_cols = [L(lay.role_col[k]) for k, _ in lay.roles]
    f_role_block = "Features!${0}$2:${1}${2}".format(role_cols[0], role_cols[-1], lay.f_last)
    feat_tot = lay.frng(lay.f_total)
    incl = 'SUMIFS({0},{1},"Included")'.format(t("E"), t("I"))
    incl_com = 'SUMIFS({0},{1},"Included",{2},"Committed")'.format(t("E"), t("I"), t("K"))
    non_shared = 'SUMIFS({0},{1},"<>SHARED")'.format(t("E"), t("A"))
    uiux_tasks = 'SUMIFS({0},{1},"UIUX",{2},"Included")'.format(t("E"), t("D"), t("I"))
    uiux_feat = 'SUMIFS({0},{1},"Included")+SUMIFS({2},{3},"SHARED",{4},"UIUX")'.format(
        lay.frng(lay.role_col["UIUX"]), lay.frng("I"), t("E"), t("A"), t("D"))
    feat_row_sum = "+".join("Features!${0}$2:${0}${1}".format(col, lay.f_last) for col in role_cols)
    com_row = "Summary!${0}${1}:${2}${1}".format(c["ms_first"], c["com_row"], c["ms_last"])
    ms_parts = []
    for row in range(int(rf), int(tot) + 1):
        rng = "Summary!${0}${2}:${1}${2}".format(c["ms_first"], c["ms_last"], row)
        ms_parts.append("ABS(SUM({0})-Summary!${1}${2})<{3}".format(rng, pc, row, EPS))
        for col, word in ((c["c_col"], "Committed"), (c["t_col"], "Tentative")):
            ms_parts.append('ABS(SUMIFS({0},{1},"{2}")-Summary!${3}${4})<{5}'.format(rng, com_row, word, col, row, EPS))
    ms_checks = ",".join(ms_parts)
    non_committed = 'SUMIFS({0},{1},"Deferred")+SUMIFS({0},{1},"Excluded")+SUMIFS({0},{1},"Pending")'.format(
        t("E"), t("I"))
    prio_sum = "+".join('COUNTIFS({0},"Included",{1},"{2}")'.format(lay.frng("I"), lay.frng("E"), p)
                        for p in PRIORITIES)
    rec_no = ('COUNTIF(Recommended!$L$2:$L${0},"No")={1}'.format(lay.r_last, lay.k) if lay.k else "TRUE")
    dup = "SUMPRODUCT(--(COUNTIFS({0},{0},{1},{1},{2},{2})>1))".format(t("A"), t("C"), t("D"))
    role_keys = "Summary!$B${0}:$B${1}".format(rf, c["role_last"])
    bad_role = "SUMPRODUCT(--(COUNTIF({0},{1})=0))".format(role_keys, t("D"))
    bad_feat = 'SUMPRODUCT(({0}<>"SHARED")*(COUNTIF({1},{0})=0))'.format(t("A"), lay.frng("A"))
    if data.get("schedule"):
        eta = '=IF({0}="{1}","PASS",IF(LEFT({0},8)="Per-role",IF({2}>0,"REVIEW","PASS"),"FAIL"))'.format(
            c["eta_method"], PENDING, c["eta_gaps"])
        eta_d = '=IF({0}>0,{0}&" role(s) have hours but no people: staff them before trusting the ETA",' \
                '"Working days come from each role\'s own hours and people, never total effort / team size")'.format(c["eta_gaps"])
    else:
        eta = '=IF(AND({0}="{1}",COUNTIF({2},"{1}")={3}),"PASS","FAIL")'.format(
            c["eta_method"], PENDING, c["eta_range"], 3 * len(lay.ms))
        eta_d = '="No schedule block: every ETA cell reads {0}"'.format(PENDING)
    rows = [
        ('=IF(AND(ABS({0}-({1}))<{3},ABS({0}-{2})<{3}),"PASS","FAIL")'.format(uiux_tasks, uiux_feat, c["uiux_summary"], EPS),
         '="Tasks "&{0}&" | Features + SHARED "&({1})&" | Summary "&{2}'.format(uiux_tasks, uiux_feat, c["uiux_summary"])),
        ('=IF(ABS(SUM({0})-{1})<{2},"PASS","FAIL")'.format(f_role_block, non_shared, EPS),
         '="Features role cells "&SUM({0})&" | Tasks (non-SHARED) "&{1}'.format(f_role_block, non_shared)),
        ('=IF(SUMPRODUCT(--(ABS({0}-({1}))>{2}))=0,"PASS","FAIL")'.format(feat_tot, feat_row_sum, EPS),
         '="Rows where Total Hours differs from the role sum: "&SUMPRODUCT(--(ABS({0}-({1}))>{2}))'.format(feat_tot, feat_row_sum, EPS)),
        ('=IF(AND({0}),"PASS","FAIL")'.format(ms_checks),
         '="Each role row: milestone columns sum to the combined total, committed columns to the Committed '
         'subtotal, tentative columns to the Tentative subtotal (an undefined milestone fails this)"'),
        (eta, eta_d),
        ('=IF(COUNTIF({0},"Included")={1},"PASS","REVIEW")'.format(lay.frng("I"), prio_sum),
         '=(COUNTIF({0},"Included")-({1}))&" included feature(s) without a P0 to P3 priority"'.format(lay.frng("I"), prio_sum)),
        ('=IF(ABS({0}-(SUM({1})-({2})))<{3},"PASS","FAIL")'.format(all_incl, t("E"), non_committed, EPS),
         '="Committed + tentative "&{0}&" = all task hours "&SUM({1})&" minus Deferred, Excluded and Pending "&({2})'.format(
             all_incl, t("E"), non_committed)),
        ('=IF(AND(ABS({0}-{1})<{3},ABS({2}-{4})<{3},{5}),"PASS","FAIL")'.format(approved, incl_com, all_incl, EPS, incl, rec_no),
         '="Recommended effort "&{0}&" h sits in the Proposed bucket only, Included in Estimate is No for every item"'.format(
             c["b_prop_rec"])),
        ('=IF({0}=0,"PASS","FAIL")'.format(dup), '={0}&" task row(s) share a Feature + Task + Role"'.format(dup)),
        ('=IF(AND(ABS({0}-{1})<{3},ABS(SUM({2})-{4})<{3}),"PASS","FAIL")'.format(approved, incl_com, feat_tot, EPS, non_shared),
         '="Committed and feature totals equal the raw sum of typed task hours: no factor is applied"'),
        ('=IF(COUNTBLANK({0})>0,"REVIEW","PASS")'.format(t("E")),
         '=COUNTBLANK({0})&" task(s) without Min Hours: counted as Unestimated, never as zero"'.format(t("E"))),
        ('=IF(AND({0}=0,{1}=0),"PASS","FAIL")'.format(bad_role, bad_feat),
         '={0}&" unknown role(s), "&{1}&" unknown feature reference(s)"'.format(bad_role, bad_feat)),
    ] + release_checks(lay)
    for i, (name, (res, det)) in enumerate(zip(CHECKS, rows), start=1):
        r = 3 + i
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=res).font = BOLD
        ws.cell(row=r, column=4, value=det)
        lay.cells["v{0}".format(i)] = "Validation!$C${0}".format(r)
    last = 3 + len(CHECKS)
    ws.cell(row=last + 2, column=2, value="Overall").font = BOLD
    ws.cell(row=last + 2, column=3, value='=IF(COUNTIF(C4:C{0},"FAIL")>0,"FAIL",IF(COUNTIF(C4:C{0},"REVIEW")>0,'
            '"REVIEW","PASS"))'.format(last)).font = BOLD
    lay.cells["v_overall"] = "Validation!$C${0}".format(last + 2)
    rng = "C4:C{0}".format(last + 2)
    for word, color in (("PASS", PASS_FILL), ("FAIL", FAIL_FILL), ("REVIEW", REVIEW_FILL)):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"{0}"'.format(word)],
                                                      fill=fill(color)))
    for col, w in zip("ABCD", (5, 58, 10, 90)):
        ws.column_dimensions[col].width = w


def build(data: dict, out_path: str | Path) -> Layout:
    """Build the five-sheet workbook and return its layout."""
    errors, _warnings = check_input(data)
    if errors:
        raise ValueError("; ".join(errors))
    lay = Layout(data)
    wb = Workbook()
    build_summary(wb, data, lay)
    build_features(wb, data, lay)
    build_tasks(wb, data, lay)
    build_recommended(wb, data, lay)
    build_validation(wb, data, lay)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return lay


def expected(data: dict) -> dict:
    """Recompute every total from the JSON alone, independent of any formula."""
    feats = {f["id"]: f for f in data["features"]}
    role_ids = [k for k, _ in roles_of(data)]
    ms_ids = [m["id"] for m in column_order(data["milestones"])]
    ms_by_id = {m["id"]: m for m in data["milestones"]}

    def scope(t: dict) -> str:
        return "Included" if t["feature"] == "SHARED" else feats.get(t["feature"], {}).get("scopeStatus", "Unknown feature")

    def ms_of(t: dict) -> str:
        return t.get("milestone") if t["feature"] == "SHARED" else feats.get(t["feature"], {}).get("milestone", "Unknown feature")

    def hrs(t: dict) -> float:
        return float(t["minHours"]) if t.get("minHours") is not None else 0.0

    def com_of(t: dict) -> str:
        m = ms_by_id.get(ms_of(t))
        return ms_commitment(m) if m else "Unknown milestone"

    def rel_of(t: dict) -> str:
        m = ms_by_id.get(ms_of(t))
        return ms_release(m) if m else "Unknown milestone"

    def upper_of(sel: list[dict]) -> float:
        return sum(float(t["maxHours"]) if t.get("maxHours") is not None else hrs(t) for t in sel)

    tasks = data["tasks"]
    grid = {(r, m): sum(hrs(t) for t in tasks if t["role"] == r and ms_of(t) == m and scope(t) == "Included")
            for r in role_ids for m in ms_ids}
    proj = {r: sum(hrs(t) for t in tasks if t["role"] == r and scope(t) == "Included") for r in role_ids}
    incl_tasks = [t for t in tasks if scope(t) == "Included"]
    com_tasks = [t for t in incl_tasks if com_of(t) == "Committed"]
    tent_tasks = [t for t in incl_tasks if com_of(t) == "Tentative"]
    proj_c = {r: sum(hrs(t) for t in com_tasks if t["role"] == r) for r in role_ids}
    proj_t = {r: sum(hrs(t) for t in tent_tasks if t["role"] == r) for r in role_ids}
    approved = sum(proj_c.values())
    all_incl_total = sum(proj.values())
    by_scope = {s: sum(hrs(t) for t in tasks if scope(t) == s) for s in ("Pending", "Deferred", "Excluded")}
    rec = sum(float(x["effortHours"]) for x in data.get("recommended", []) if x.get("effortHours") is not None)
    upper = upper_of(com_tasks)
    out: dict[str, Any] = {
        "approved": approved, "all_incl": all_incl_total, "tent_total": sum(proj_t.values()),
        "b_approved": approved, "u_approved": upper, "b_tent_total": sum(proj_t.values()),
        "u_tent_total": upper_of(tent_tasks), "proj_c": proj_c, "proj_t": proj_t,
        "b_prop_feat": by_scope["Pending"], "b_prop_rec": rec,
        "b_prop_total": by_scope["Pending"] + rec, "b_deferred": by_scope["Deferred"],
        "b_excluded": by_scope["Excluded"], "b_upper": upper,
        "b_unest_tasks": float(sum(1 for t in tasks if t.get("minHours") is None)),
        "b_unest_feats": float(sum(1 for f in data["features"] if not any(t["feature"] == f["id"] for t in tasks))),
        "grid": grid, "proj": proj,
    }
    for p in PRIORITIES:
        out["cnt_" + p] = float(sum(1 for f in data["features"] if f.get("priority") == p))
    for rname in RECOMMENDATIONS:
        out["cnt_" + rname] = float(sum(1 for f in data["features"] if f.get("recommendation") == rname))
    for i, rel in enumerate(tentative_releases(data["milestones"])):
        sel = [t for t in tent_tasks if rel_of(t) == rel]
        out["b_tent_{0}".format(i)] = sum(hrs(t) for t in sel)
        out["u_tent_{0}".format(i)] = upper_of(sel)
    unresolved = 0
    for f in data["features"]:
        ds = deps_list(f.get("dependsOn"))
        if ds:
            ok = all(feats.get(d, {}).get("scopeStatus") == "Included" for d in ds)
            out["dep_" + f["id"]] = "Resolved" if ok else "Unresolved"
            unresolved += 0 if ok else 1
    out["dep_unresolved"] = float(unresolved)

    gaps = 0
    sched = data.get("schedule")
    if sched:
        hpd = float(data["project"].get("hoursPerDay", 8))
        people = sched.get("people", {})
        start = date_to_serial(dt.date.fromisoformat(sched["startDate"]))
        cur = workday(start - 1, 1)
        for m in ms_ids:
            ratios = [grid[(r, m)] / (people.get(r, 0) * hpd) for r in role_ids if people.get(r, 0) > 0]
            gaps += sum(1 for r in role_ids if grid[(r, m)] > 0 and not people.get(r, 0) > 0)
            days = float(math.ceil(max([0.0] + ratios)))
            end = cur if days == 0 else workday(cur, days - 1)
            out["eta_days_" + m], out["eta_end_" + m] = days, end
            cur = workday(end, 1)
        out["eta_gaps"] = float(gaps)
        com_ms = [m for m in ms_ids if ms_commitment(ms_by_id[m]) == "Committed"]
        tent_ms = [m for m in ms_ids if ms_commitment(ms_by_id[m]) != "Committed"]
        out["eta_delivery"] = out["eta_end_" + com_ms[-1]] if com_ms else "None committed"
        out["eta_tent_end"] = out["eta_end_" + tent_ms[-1]] if tent_ms else "No tentative phases"

    known_feat = all(t["feature"] == "SHARED" or t["feature"] in feats for t in tasks)
    known_role = all(t["role"] in role_ids for t in tasks)
    keys = [(str(t["feature"]).lower(), str(t["task"]).lower(), str(t["role"]).lower()) for t in tasks]
    all_hours = sum(hrs(t) for t in tasks)
    incl_all = sum(hrs(t) for t in tasks if scope(t) == "Included")
    non_shared = sum(hrs(t) for t in tasks if t["feature"] != "SHARED")
    feat_cells = sum(hrs(t) for t in tasks if t["feature"] in feats and t["role"] in role_ids)
    ms_ok = all(abs(sum(grid[(r, m)] for m in ms_ids) - proj[r]) < EPS for r in role_ids)
    uiux_ok = abs(proj["UIUX"] - (sum(hrs(t) for t in tasks if t["role"] == "UIUX" and t["feature"] in feats
                                      and scope(t) == "Included")
                                  + sum(hrs(t) for t in tasks if t["feature"] == "SHARED" and t["role"] == "UIUX"))) < EPS
    incl_feats = [f for f in data["features"] if f.get("scopeStatus") == "Included"]

    def f_com(f: dict) -> str:
        m = ms_by_id.get(f.get("milestone"))
        return ms_commitment(m) if m else "Unknown milestone"

    leak = [f for f in incl_feats if f_com(f) == "Committed" and f.get("scopeDecision") == "Tentative"]
    low = [f for f in incl_feats if f_com(f) == "Committed" and f.get("priority") not in ("P0", "P1")
           and not f.get("reasoning")]
    defaulted = [m for m in data["milestones"] if not m.get("release") or not m.get("commitment")
                 or m.get("commitment") not in COMMITMENTS]
    v = {
        1: "PASS" if uiux_ok else "FAIL",
        2: "PASS" if abs(feat_cells - non_shared) < EPS else "FAIL",
        3: "PASS",
        4: "PASS" if ms_ok else "FAIL",
        5: ("REVIEW" if gaps else "PASS"),
        6: "PASS" if all(f.get("priority") in PRIORITIES for f in incl_feats) else "REVIEW",
        7: "PASS" if abs(all_incl_total - (all_hours - sum(by_scope.values()))) < EPS else "FAIL",
        8: "PASS" if abs(all_incl_total - incl_all) < EPS else "FAIL",
        9: "PASS" if len(keys) == len(set(keys)) else "FAIL",
        10: "PASS" if abs(all_incl_total - incl_all) < EPS and abs(feat_cells - non_shared) < EPS else "FAIL",
        11: "REVIEW" if any(t.get("minHours") is None for t in tasks) else "PASS",
        12: "PASS" if known_feat and known_role else "FAIL",
        13: "FAIL" if leak else "PASS",
        14: "REVIEW" if low else "PASS",
        15: "REVIEW" if defaulted else "PASS",
    }
    out["validation"] = v
    vals = list(v.values())
    out["v_overall"] = "FAIL" if "FAIL" in vals else ("REVIEW" if "REVIEW" in vals else "PASS")
    return out


def _get(ev: Evaluator, addr: str) -> Any:
    """Evaluate one 'Sheet!$A$1' address."""
    sheet, cell = addr.split("!")
    return ev.value(sheet, cell.replace("$", ""))


def verify(path: str | Path, data: dict, lay: Layout) -> list[str]:
    """Evaluate the workbook's formulas in Python and compare with expected(); raise on mismatch."""
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Features", "Tasks", "Recommended", "Validation"], wb.sheetnames
    ev = Evaluator(wb)
    exp = expected(data)
    problems: list[str] = []

    def cmp(label: str, got: Any, want: Any) -> None:
        if isinstance(want, float):
            if not isinstance(got, (int, float)) or abs(float(got) - want) > EPS:
                problems.append("{0}: formula gives {1!r}, expected {2!r}".format(label, got, want))
        elif got != want:
            problems.append("{0}: formula gives {1!r}, expected {2!r}".format(label, got, want))

    cmp("Committed", _get(ev, lay.cells["approved"]), exp["approved"])
    cmp("Tentative", _get(ev, lay.cells["tent_total"]), exp["tent_total"])
    cmp("Committed + tentative", _get(ev, lay.cells["all_incl"]), exp["all_incl"])
    for key in [k for k in exp if k.startswith(("b_tent_", "u_"))] + ["b_approved"]:
        cmp(key, _get(ev, lay.cells[key]), exp[key])
    for key in ("b_prop_feat", "b_prop_rec", "b_prop_total", "b_deferred", "b_excluded", "b_upper",
                "b_unest_tasks", "b_unest_feats", "dep_unresolved"):
        cmp(key, _get(ev, lay.cells[key]), exp[key])
    for key in [k for k in exp if k.startswith(("cnt_", "dep_F", "eta_"))]:
        cmp(key, _get(ev, lay.cells[key]), exp[key])
    rf = int(lay.cells["role_first"])
    for i, (role, _) in enumerate(lay.roles):
        for m in lay.ms:
            cmp("{0}/{1}".format(role, m["id"]), ev.value("Summary", "{0}{1}".format(L(lay.ms_col[m["id"]]), rf + i)),
                exp["grid"][(role, m["id"])])
        for col, key in (("proj_col", "proj"), ("c_col", "proj_c"), ("t_col", "proj_t")):
            cmp("{0}/{1}".format(role, key), ev.value("Summary", "{0}{1}".format(lay.cells[col], rf + i)),
                exp[key][role])
    for i in range(1, len(CHECKS) + 1):
        cmp("Validation check {0}".format(i), _get(ev, lay.cells["v{0}".format(i)]), exp["validation"][i])
    cmp("Validation overall", _get(ev, lay.cells["v_overall"]), exp["v_overall"])
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    got = ev.cell(ws.title, cell.coordinate)
                    if type(got).__name__ == "XLError":
                        problems.append("{0}!{1} evaluates to an error".format(ws.title, cell.coordinate))
                    if ws.title != "Tasks" and _has_factor(cell.value):
                        problems.append("{0}!{1} multiplies hours by a factor".format(ws.title, cell.coordinate))
    if problems:
        raise AssertionError("verify failed:\n  " + "\n  ".join(problems))
    return ["verify: {0} cells evaluated, totals match the JSON recomputation".format(len(ev.cache)),
            "Committed {0:.1f} h (upper {1:.1f}) | Tentative {2:.1f} h (upper {3:.1f}) | Validation {4}".format(
                exp["approved"], exp["b_upper"], exp["tent_total"], exp["u_tent_total"], exp["v_overall"])]


def _has_factor(formula: str) -> bool:
    """True when a formula multiplies by a numeric literal (a margin-style factor)."""
    import re
    literal_after = re.search(r"\*\s*\d+(?:\.\d+)?(?![\d$A-Z])", formula)
    literal_before = re.search(r"(?<![A-Z$\d])\d+(?:\.\d+)?\s*\*", formula)
    return bool(literal_after or literal_before)


def run_selftest() -> None:
    """Build both bundled examples plus five negative cases and verify every formula."""
    here = Path(__file__).resolve().parent.parent / "examples"
    tmp = Path(__file__).resolve().parent.parent / ".selftest"
    tmp.mkdir(exist_ok=True)
    try:
        for name in ("estimation-example-v3.json", "estimation-example-v3-no-schedule.json"):
            data = json.loads((here / name).read_text(encoding="utf-8"))
            lay = build(data, tmp / (name + ".xlsx"))
            for line in verify(tmp / (name + ".xlsx"), data, lay):
                print("  [{0}] {1}".format(name, line))
        base = json.loads((here / "estimation-example-v3.json").read_text(encoding="utf-8"))
        dup = copy.deepcopy(base)
        dup["tasks"].append(copy.deepcopy(next(t for t in dup["tasks"] if t.get("minHours") is not None)))
        bad = copy.deepcopy(base)
        bad["tasks"].append({"feature": "F99", "task": "Orphan", "role": "Designer", "minHours": 4,
                             "confidence": "Low", "assumption": "negative test"})
        leak = copy.deepcopy(base)
        tent_feat = next(f for f in leak["features"] if f.get("scopeDecision") == "Tentative"
                         and f.get("scopeStatus") == "Included")
        tent_feat["milestone"] = next(m["id"] for m in leak["milestones"] if m.get("commitment") == "Committed")
        low = copy.deepcopy(base)
        low_feat = next(f for f in low["features"] if f.get("scopeStatus") == "Included"
                        and f.get("scopeDecision") == "Committed" and f.get("priority") == "P1")
        low_feat["priority"], low_feat["reasoning"] = "P3", ""
        default = json.loads((here / "estimation-example-v3-no-schedule.json").read_text(encoding="utf-8"))
        cases = (("duplicate task", dup, 9, "FAIL"), ("unknown role and feature", bad, 12, "FAIL"),
                 ("tentative feature " + tent_feat["id"] + " placed in a committed milestone", leak, 13, "FAIL"),
                 ("committed P3 feature " + low_feat["id"] + " without reasoning", low, 14, "REVIEW"),
                 ("milestones without release or commitment (defaulted)", default, 15, "REVIEW"))
        for label, case, check, want in cases:
            lay = build(case, tmp / "neg.xlsx")
            verify(tmp / "neg.xlsx", case, lay)
            got = _get(Evaluator(load_workbook(tmp / "neg.xlsx")), lay.cells["v{0}".format(check)])
            assert got == want, "{0}: check {1} should be {2}, got {3}".format(label, check, want, got)
            print("  [negative] {0}: check {1} = {2} as expected".format(label, check, want))
    finally:
        for f in tmp.glob("*"):
            f.unlink()
        tmp.rmdir()
    print("build_workbook selftest: PASS")


def main(argv: list[str] | None = None) -> int:
    """CLI entry point."""
    ap = argparse.ArgumentParser(description="Build the simple estimation workbook (schemaVersion 3).")
    ap.add_argument("input", nargs="?", help="estimation JSON (schemaVersion 3)")
    ap.add_argument("output", nargs="?", help="output .xlsx path (default: next to the JSON)")
    ap.add_argument("--selftest", action="store_true", help="build the bundled examples and verify them")
    args = ap.parse_args(argv)
    if args.selftest:
        run_selftest()
        return 0
    if not args.input:
        ap.error("input JSON is required unless --selftest is given")
    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    errors, warnings = check_input(data)
    if errors:
        print("ERROR: " + "; ".join(errors), file=sys.stderr)
        return 2
    for w in warnings:
        print("  warning: " + w)
    out = args.output or str(Path(args.input).with_suffix(".xlsx"))
    lay = build(data, out)
    for line in verify(out, data, lay):
        print("  " + line)
    print("Wrote {0}".format(out))
    print("  Sheets: Summary | Features | Tasks | Recommended | Validation")
    print("  Open it in Excel (or LibreOffice) once so every formula recalculates before sharing.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
