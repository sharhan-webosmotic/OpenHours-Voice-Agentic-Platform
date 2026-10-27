"""PPTX coverage inventory: one unit per slide.

Records text, tables, pictures, charts, shape diagrams (groups, connectors,
SmartArt), embedded objects, speaker notes and the hidden flag. Connector
relationships are recovered from the connection ids when present, else from
geometry (the shape whose box is nearest each connector end).
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from inv_common import chart_info, flag, new_unit, snippet
from office_render import normalize_image, render_pptx_slides

EDGE_TOLERANCE_EMU = 457200


def _uri(shape) -> str:
    """Return the graphicData uri of a graphic frame, or an empty string."""
    found = shape._element.xpath(".//a:graphicData/@uri")
    return found[0] if found else ""


def _smartart_text(slide, shape) -> str:
    """Return the text of a SmartArt diagram stored in its diagram data part."""
    from lxml import etree

    xml = etree.tostring(shape._element).decode("utf-8", "ignore")
    texts = []
    for rid in re.findall(r'r:dm="([^"]+)"', xml):
        if rid in slide.part.rels:
            blob = slide.part.related_part(rid).blob.decode("utf-8", "ignore")
            texts += re.findall(r"<a:t>([^<]*)</a:t>", blob)
    return " | ".join(t for t in texts if t.strip())


def _distance(px: int, py: int, box: tuple[int, int, int, int]) -> float:
    """Return the distance from a point to an axis-aligned box (0 when inside)."""
    left, top, width, height = box
    dx = max(left - px, 0, px - (left + width))
    dy = max(top - py, 0, py - (top + height))
    return (dx * dx + dy * dy) ** 0.5


def _nearest(px: int, py: int, boxes: list[dict]) -> str:
    """Return the text of the labelled shape nearest a point, within tolerance."""
    best, best_d = "", float(EDGE_TOLERANCE_EMU)
    for box in boxes:
        d = _distance(px, py, box["box"])
        if d <= best_d:
            best, best_d = box["text"], d
    return best


def _walk(slide, shapes, unit: dict, state: dict) -> None:
    """Recursively record every shape on a slide into the unit and walk state."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.shapes.connector import Connector

    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            unit["diagram"]["groups"] += 1
            _walk(slide, sh.shapes, unit, state)
            continue
        if isinstance(sh, Connector):
            unit["diagram"]["connectors"] += 1
            state["connectors"].append(sh)
            continue
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            text = sh.text_frame.text.strip()
            unit["text_chars"] += len(text)
            if sh.width is not None and sh.left is not None:
                state["boxes"].append(
                    {"id": sh.shape_id, "text": text, "box": (sh.left, sh.top, sh.width, sh.height)}
                )
        if hasattr(sh, "image"):
            try:
                state["media"].append((sh.image.blob, sh.image.ext))
                unit["pictures"] += 1
            except (AttributeError, ValueError, KeyError):
                pass
        if getattr(sh, "has_table", False):
            rows, cols = len(sh.table.rows), len(sh.table.columns)
            unit["tables"].append({"rows": rows, "cols": cols})
            unit["text_chars"] += sum(len(c.text) for r in sh.table.rows for c in r.cells)
        if getattr(sh, "has_chart", False):
            unit["charts"].append(chart_info(sh.chart_part.blob))
        uri = _uri(sh) if sh._element.tag.endswith("}graphicFrame") else ""
        if "diagram" in uri:
            unit["diagram"]["smartart"] += 1
            text = _smartart_text(slide, sh)
            if text:
                state["smartart_text"].append(text)
        elif "ole" in uri.lower():
            unit["embedded_objects"] += 1


def _edges(state: dict) -> list[list[str]]:
    """Resolve each connector to a [from, to] pair of shape labels."""
    by_id = {b["id"]: b["text"] for b in state["boxes"]}
    edges = []
    for cx in state["connectors"]:
        st = cx._element.xpath(".//a:stCxn/@id")
        en = cx._element.xpath(".//a:endCxn/@id")
        src = by_id.get(int(st[0]), "") if st else ""
        dst = by_id.get(int(en[0]), "") if en else ""
        src = src or _nearest(cx.begin_x, cx.begin_y, state["boxes"])
        dst = dst or _nearest(cx.end_x, cx.end_y, state["boxes"])
        edges.append([src or "?", dst or "?", _direction(cx)])
    return edges


def _direction(cx) -> str:
    """Return the edge operator from the connector's arrowheads.

    In DrawingML ``a:headEnd`` decorates the line start and ``a:tailEnd`` the
    line end. No arrowhead means the slide shows no direction, so the edge is
    undirected rather than assumed to run begin-to-end.
    """
    def has_arrow(tag: str) -> bool:
        types = cx._element.xpath(f".//a:ln/a:{tag}/@type")
        return bool(types) and types[0] != "none"

    start, end = has_arrow("headEnd"), has_arrow("tailEnd")
    if start and end:
        return "<->"
    if end:
        return "->"
    if start:
        return "<-"
    return "--"


def _comments(src: Path) -> list[dict]:
    """Return slide comment texts from legacy and modern comment parts."""
    out = []
    with zipfile.ZipFile(src) as z:
        for name in z.namelist():
            if name.startswith("ppt/comments/") and name.endswith(".xml"):
                xml = z.read(name).decode("utf-8", "ignore")
                for text in re.findall(r"<p(?:188)?:text>([^<]*)</p(?:188)?:text>", xml):
                    out.append({"kind": "comment", "text": text, "probe": snippet(text)})
                for text in re.findall(r"<a:t>([^<]*)</a:t>", xml):
                    out.append({"kind": "comment", "text": text, "probe": snippet(text)})
    return [c for c in out if c["text"].strip()]


def build(src: Path, ctx: dict) -> dict:
    """Inventory a PPTX and render flagged slides. Returns {units, document, render}."""
    from pptx import Presentation

    prs = Presentation(str(src))
    units, media_by_unit = [], {}
    for idx, slide in enumerate(prs.slides, 1):
        unit = new_unit(f"slide-{idx}", idx, f"Slide {idx}")
        state = {"boxes": [], "connectors": [], "media": [], "smartart_text": []}
        _walk(slide, slide.shapes, unit, state)
        unit["diagram"]["edges"] = _edges(state)
        unit["smartart_text"] = state["smartart_text"]
        unit["hidden"] = slide._element.get("show") == "0"
        unit["expect_md_tables"] = len(unit["tables"])
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            unit["notes"] = notes
            if notes:
                unit["probes"].append(snippet(notes))
        for count, reason in (
            (unit["pictures"], "picture"),
            (len(unit["charts"]), "chart"),
            (unit["diagram"]["connectors"], "connector"),
            (unit["diagram"]["groups"], "group"),
            (unit["diagram"]["smartart"], "smartart"),
            (unit["embedded_objects"], "embedded-object"),
        ):
            if count:
                flag(unit, reason)
        units.append(unit)
        media_by_unit[unit["id"]] = state["media"]

    render = {"engine": "none", "note": ""}
    flagged = {u["index"]: ctx["assets_dir"] / f"slide-{u['index']:03d}.png" for u in units if u["needs_visual"]}
    if flagged and ctx["do_render"]:
        err = render_pptx_slides(src, flagged, ctx["temp_dir"])
        render = {"engine": "powerpoint-com", "note": ""} if not err else {"engine": "media-fallback", "note": f"slide render unavailable: {err}"}
    for unit in units:
        if not unit["needs_visual"]:
            continue
        png = flagged[unit["index"]]
        if png.exists():
            unit["images"] = [f"{ctx['rel']}/{png.name}"]
            continue
        for k, (blob, _ext) in enumerate(media_by_unit[unit["id"]], 1):
            dest = ctx["assets_dir"] / f"slide-{unit['index']:03d}-media-{k:02d}.png"
            if normalize_image(blob, dest) is None:
                unit["images"].append(f"{ctx['rel']}/{dest.name}")
        if not unit["images"]:
            unit["render_note"] = "slide render unavailable"
    document = {"comments": _comments(src), "hidden_units": [u["id"] for u in units if u["hidden"]]}
    return {"units": units, "document": document, "render": render}
