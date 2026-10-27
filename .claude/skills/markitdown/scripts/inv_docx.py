"""DOCX coverage inventory: a body unit plus one unit per figure.

A figure is any drawing, VML picture, or embedded object in the main body, in
document order. Pictures are flagged and their media is exported. Charts,
SmartArt, shapes and groups are flagged and rendered by exporting the document
to PDF through Word COM and rasterizing the page that holds the figure.
Headers, footers, comments, footnotes, endnotes and text boxes are recorded as
document-level parts with a matchable probe.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from inv_common import chart_info, flag, new_unit, snippet
from office_render import export_docx_pdf, normalize_image, pdf_page_count, render_pdf_pages

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
FIGURE_TAGS = (f"{W}drawing", f"{W}pict", f"{W}object")
MAX_FALLBACK_PAGES = 5


def _ancestors(el):
    """Yield every ancestor of an lxml element."""
    parent = el.getparent()
    while parent is not None:
        yield parent
        parent = parent.getparent()


def _own_text(p) -> str:
    """Return a paragraph's run text, excluding text nested in text boxes."""
    return "".join(
        t.text or "" for t in p.iter(f"{W}t") if not any(a.tag == f"{W}txbxContent" for a in _ancestors(t))
    ).strip()


def _box_text(el) -> str:
    """Return the text inside any text box nested in a figure element."""
    return " ".join(
        "".join(t.text or "" for t in p.iter(f"{W}t")).strip()
        for box in el.iter(f"{W}txbxContent")
        for p in box.iter(f"{W}p")
    ).strip()


def _classify(el, part) -> dict:
    """Classify one figure element into kind, media blob, chart, and box text."""
    xml = el.xpath("string(.//*[local-name()='graphicData']/@uri)") if el.tag == f"{W}drawing" else ""
    fig = {"kind": "shape", "media": None, "chart": None, "box_text": _box_text(el)}
    blips = el.xpath(".//*[local-name()='blip']/@*[local-name()='embed'] | .//*[local-name()='imagedata']/@*[local-name()='id']")
    if el.tag == f"{W}object":
        fig["kind"] = "embedded-object"
    elif "picture" in xml or (el.tag == f"{W}pict" and blips):
        fig["kind"] = "picture"
    elif "chart" in xml:
        fig["kind"] = "chart"
        rid = el.xpath("string(.//*[local-name()='chart']/@*[local-name()='id'])")
        if rid in part.rels:
            fig["chart"] = chart_info(part.related_parts[rid].blob)
    elif "diagram" in xml:
        fig["kind"] = "smartart"
    elif "wordprocessingGroup" in xml or "wordprocessingCanvas" in xml:
        fig["kind"] = "group"
    elif fig["box_text"]:
        fig["kind"] = "textbox"
    if blips and blips[0] in part.rels:
        fig["media"] = part.related_parts[blips[0]].blob
    return fig


def _figures(doc) -> tuple[list[dict], int, int]:
    """Walk the body in order. Returns (figures, body_text_chars, top_tables)."""
    body = doc.element.body
    figures, chars = [], 0
    last_text = ""
    for p in body.iter(f"{W}p"):
        if any(a.tag == f"{W}txbxContent" for a in _ancestors(p)):
            continue
        text = _own_text(p)
        chars += len(text)
        for el in p.iter(*FIGURE_TAGS):
            if any(a.tag == MC_FALLBACK or a.tag in FIGURE_TAGS for a in _ancestors(el)):
                continue
            fig = _classify(el, doc.part)
            fig["anchor"] = (text or last_text)[-60:]
            figures.append(fig)
        if text:
            last_text = text
    return figures, chars, len(body.findall(f"{W}tbl"))


def _story_text(story) -> str:
    """Return the paragraph and table text of a header or footer story."""
    parts = [p.text for p in story.paragraphs if p.text.strip()]
    parts += [c.text for t in story.tables for r in t.rows for c in r.cells if c.text.strip()]
    return " ".join(parts).strip()


def _headers_footers(doc) -> list[dict]:
    """Return distinct header and footer texts with section and type labels."""
    found, seen = [], set()
    for i, sec in enumerate(doc.sections, 1):
        for kind, attr in (
            ("Header", "header"), ("First-page header", "first_page_header"), ("Even-page header", "even_page_header"),
            ("Footer", "footer"), ("First-page footer", "first_page_footer"), ("Even-page footer", "even_page_footer"),
        ):
            story = getattr(sec, attr)
            if story.is_linked_to_previous:
                continue
            text = _story_text(story)
            if text and (kind, text) not in seen:
                seen.add((kind, text))
                found.append({"kind": "footer" if "ooter" in kind else "header", "label": f"{kind} (section {i})", "text": text, "probe": snippet(text)})
    return found


def _notes_parts(src: Path) -> list[dict]:
    """Return comments, footnotes and endnotes read straight from the package XML."""
    out = []
    with zipfile.ZipFile(src) as z:
        names = set(z.namelist())
        for part, tag, kind in (
            ("word/comments.xml", "comment", "comment"),
            ("word/footnotes.xml", "footnote", "footnote"),
            ("word/endnotes.xml", "endnote", "endnote"),
        ):
            if part not in names:
                continue
            xml = z.read(part).decode("utf-8", "ignore")
            for attrs, inner in re.findall(rf"<w:{tag}\b([^>]*)>(.*?)</w:{tag}>", xml, re.S):
                if 'w:type="separator"' in attrs or 'w:type="continuationSeparator"' in attrs:
                    continue
                text = "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", inner)).strip()
                author = re.search(r'w:author="([^"]*)"', attrs)
                if text:
                    label = f"Comment by {author.group(1)}" if author else kind.capitalize()
                    out.append({"kind": kind, "label": label, "text": text, "probe": snippet(text)})
    return out


def _page_map(figures: list[dict], com: dict | None) -> list[int]:
    """Map each figure to a 1-based page using Word's anchor positions."""
    anchors = sorted((f.get("start", 0), f.get("page", 1)) for f in (com or {}).get("figures", []) or [])
    if anchors and len(anchors) == len(figures):
        return [pg for _start, pg in anchors]
    return [0] * len(figures)


def build(src: Path, ctx: dict) -> dict:
    """Inventory a DOCX and render flagged figures. Returns {units, document, render}."""
    import docx

    doc = docx.Document(str(src))
    figures, chars, tables = _figures(doc)
    body = new_unit("body", 0, "Document body")
    body["text_chars"] = chars
    body["tables"] = [{"rows": len(t.rows), "cols": len(t.columns)} for t in doc.tables]
    body["expect_md_tables"] = tables
    units = [body]
    for i, fig in enumerate(figures, 1):
        unit = new_unit(f"figure-{i}", i, f"Figure {i} ({fig['kind']})")
        unit["kind"], unit["anchor"] = fig["kind"], fig["anchor"]
        if fig["chart"]:
            unit["charts"].append(fig["chart"])
        if fig["kind"] == "picture":
            unit["pictures"] = 1
        if fig["kind"] != "textbox":
            flag(unit, {"picture": "picture", "chart": "chart", "smartart": "smartart", "group": "group",
                        "embedded-object": "embedded-object"}.get(fig["kind"], "shape"))
        units.append(unit)

    render = {"engine": "none", "note": ""}
    need_page = [u for u, f in zip(units[1:], figures) if u["needs_visual"] and not f["media"]]
    pdf, com, pages = None, None, 0
    if [u for u in units if u["needs_visual"]] and ctx["do_render"]:
        pdf, com, err = export_docx_pdf(src, ctx["temp_dir"])
        render = {"engine": "word-com+pypdfium2", "note": ""} if pdf else {"engine": "media-fallback", "note": f"page render unavailable: {err}"}
        pages = pdf_page_count(pdf) if pdf else 0
    page_of = _page_map(figures, com)
    for unit, fig, page in zip(units[1:], figures, page_of):
        unit["page"] = page or None
        if not unit["needs_visual"]:
            continue
        if fig["media"] is not None:
            dest = ctx["assets_dir"] / f"figure-{unit['index']:03d}.png"
            if normalize_image(fig["media"], dest) is None:
                unit["images"].append(f"{ctx['rel']}/{dest.name}")
        if unit in need_page and pdf:
            targets = [page] if page else list(range(1, min(pages, MAX_FALLBACK_PAGES) + 1))
            dests = {n: ctx["assets_dir"] / f"page-{n:03d}.png" for n in targets}
            render_pdf_pages(pdf, dests)
            unit["images"] += [f"{ctx['rel']}/{d.name}" for d in dests.values() if d.exists()]
        if not unit["images"]:
            unit["render_note"] = "page render unavailable"

    text_boxes = [
        {"kind": "textbox", "label": f"Text box (figure {i})", "text": f["box_text"], "probe": snippet(f["box_text"])}
        for i, f in enumerate(figures, 1) if f["box_text"]
    ]
    document = {
        "parts": _headers_footers(doc) + _notes_parts(src) + text_boxes,
        "embedded_objects": sum(1 for f in figures if f["kind"] == "embedded-object"),
        "pages": pages or None,
    }
    return {"units": units, "document": document, "render": render}
